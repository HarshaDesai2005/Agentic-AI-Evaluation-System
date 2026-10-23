"""Create synthetic PDF/DOCX fixtures for tests, demo and evaluation."""

from __future__ import annotations

import json
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

import config

DEMO = config.SAMPLE_DIR / "demo"
EVAL = config.SAMPLE_DIR / "eval"
FIXTURES = config.ROOT_DIR / "tests" / "fixtures"


SYLLABUS = """DIGITAL LOGIC DESIGN
Subject Code: CS301
Official Syllabus

UNIT 1: Number Systems and Codes
Binary, octal, hexadecimal conversions. Complements. Binary codes. Error detecting codes.

UNIT 2: Boolean Algebra and Logic Gates
Boolean theorems. De Morgan's theorems. Canonical forms. Logic gates. Simplification using Boolean algebra.

UNIT 3: Combinational Logic
Karnaugh maps (up to 4 variables). Adders and subtractors. Multiplexers. Encoders. Decoders. Comparators.

UNIT 4: Sequential Logic
Latches. Flip-flops. Registers. Counters. SRAM and DRAM memory comparison at the logic-design level.

UNIT 5: HDL Introduction
Brief introduction to Verilog modelling of combinational circuits.

This syllabus does not include quantum computing, quantum annealing, or machine learning.
"""

TEXTBOOK = """Chapter 2 Boolean Algebra
De Morgan's first theorem states that the complement of a product is equal to the sum of the complements:
(A · B)' = A' + B'. The second theorem states that the complement of a sum is equal to the product of the complements:
(A + B)' = A' · B'.

Chapter 3 Number Systems
To convert a binary number to decimal, expand as a sum of powers of two.
Example: 1011012 = 1*32 + 0*16 + 1*8 + 1*4 + 0*2 + 1*1 = 45 decimal.

Chapter 4 Combinational Circuits
A 4-bit ripple carry adder is built from four full adders. The carry of stage i is connected to stage i+1.
Each full adder computes sum = A xor B xor Cin and carry = majority(A, B, Cin).

A 4-variable Karnaugh map has 16 cells. Adjacent 1s are grouped in powers of two to obtain a simplified SOP expression.

Chapter 6 Memory
SRAM uses latches and is faster. DRAM uses capacitors, needs refresh, and is denser. SRAM is used for cache; DRAM for main memory.
"""

NOTES = """Lecture notes
Full adder truth table: inputs A, B, Cin; outputs Sum and Cout.
FPGA architecture is only mentioned as optional reading and is not a core assessment topic in CS301.
"""

PREVIOUS = """Previous approved scheme
Q: State De Morgan's theorems. Total 5 marks.
2 marks for first theorem with equation.
2 marks for second theorem with equation.
1 mark for a brief explanation or example.
"""

PAPER = """CS301 Digital Logic Design
End Semester Examination
Maximum Marks: 48
Time: 3 Hours

Instructions: Answer all questions. Internal choice is provided in Q5.

PART A

Q1 (a) State and explain De Morgan's theorems with equations. (5 marks)
(b) Convert the binary number 101101 to decimal. Show the calculation. (5 marks)

Q2 Design a 4-bit ripple carry adder using full adders. Explain the carry propagation. (10 marks)

Q3 Explain quantum annealing optimisation and how it solves SAT problems. (10 marks)

Q4 Compare SRAM and DRAM with respect to storage mechanism, speed, density and typical use. (8 marks)

Q5 Answer any one of the following.
(a) Minimise the Boolean function F(A,B,C,D)=Σm(0,1,2,5,8,9,10) using a 4-variable K-map. (10 marks)
OR
(b) Explain FPGA architecture in detail, including CLBs, routing and I/O blocks. (10 marks)
"""

EVAL_PAPER = """CS301 Makeup Paper
Maximum Marks: 20

Q1 Define a full adder and write its sum and carry expressions. (5 marks)
Q2 Convert 1101 binary to decimal. (5 marks)
Q3 Describe convolutional neural networks for image classification. (10 marks)
"""


def _pdf(path: Path, text: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    c = canvas.Canvas(str(path), pagesize=A4)
    width, height = A4
    y = height - 50
    c.setFont("Times-Roman", 11)
    for line in text.splitlines():
        if y < 60 or line.startswith("Chapter"):
            if line.startswith("Chapter") and y < height - 50:
                c.showPage()
                y = height - 50
            elif y < 60:
                c.showPage()
                y = height - 50
            c.setFont("Times-Roman", 11)
        c.drawString(50, y, line[:110])
        y -= 16
    c.save()


def _docx(path: Path, text: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    for line in text.splitlines():
        if line.isupper() or line.startswith("UNIT") or line.startswith("Chapter"):
            doc.add_heading(line, level=1)
        else:
            doc.add_paragraph(line)
    doc.save(str(path))


def write_all() -> dict[str, str]:
    for folder in (DEMO, EVAL, FIXTURES):
        folder.mkdir(parents=True, exist_ok=True)
    _pdf(DEMO / "question_paper.pdf", PAPER)
    _docx(DEMO / "syllabus.docx", SYLLABUS)
    _pdf(DEMO / "textbook.pdf", TEXTBOOK)
    _docx(DEMO / "notes.docx", NOTES)
    _docx(DEMO / "previous_scheme.docx", PREVIOUS)

    _pdf(EVAL / "question_paper.pdf", EVAL_PAPER)
    _docx(EVAL / "syllabus.docx", SYLLABUS)
    _pdf(EVAL / "textbook.pdf", TEXTBOOK)

    _pdf(FIXTURES / "question_paper.pdf", PAPER)
    _docx(FIXTURES / "syllabus.docx", SYLLABUS)
    _pdf(FIXTURES / "textbook.pdf", TEXTBOOK)
    _docx(FIXTURES / "notes.docx", NOTES)
    _pdf(FIXTURES / "empty.pdf", "   \n\n")
    _docx(FIXTURES / "simple.docx", "Heading One\nBody text on logic gates and De Morgan.")

    reference = {
        "subject": "Digital Logic Design",
        "code": "CS301",
        "questions": [
            {
                "id": "Q1a",
                "expect_syllabus": "IN",
                "contains": "De Morgan",
                "marks": 5,
            },
            {
                "id": "Q1b",
                "expect_syllabus": "IN",
                "contains": "101101",
                "marks": 5,
                "numeric_answer": "45",
            },
            {
                "id": "Q2",
                "expect_syllabus": "IN",
                "contains": "ripple",
                "marks": 10,
            },
            {
                "id": "Q3",
                "expect_syllabus": "OUT",
                "contains": "quantum annealing",
                "marks": 10,
            },
            {
                "id": "Q4",
                "expect_syllabus": "IN",
                "contains": "SRAM",
                "marks": 8,
            },
            {
                "id": "Q5a",
                "expect_syllabus": "IN",
                "contains": "K-map",
                "marks": 10,
            },
            {
                "id": "Q5b",
                "expect_syllabus": "PARTIAL",
                "contains": "FPGA",
                "marks": 10,
            },
        ],
    }
    (DEMO / "reference.json").write_text(json.dumps(reference, indent=2), encoding="utf-8")
    eval_ref = {
        "questions": [
            {"id": "Q1", "expect_syllabus": "IN", "contains": "full adder", "marks": 5},
            {"id": "Q2", "expect_syllabus": "IN", "contains": "1101", "marks": 5, "numeric_answer": "13"},
            {"id": "Q3", "expect_syllabus": "OUT", "contains": "neural", "marks": 10},
        ]
    }
    (EVAL / "reference.json").write_text(json.dumps(eval_ref, indent=2), encoding="utf-8")
    return {"demo": str(DEMO), "eval": str(EVAL), "fixtures": str(FIXTURES)}


if __name__ == "__main__":
    print(write_all())
