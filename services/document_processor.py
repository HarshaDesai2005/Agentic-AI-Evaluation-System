"""PDF/DOCX/image ingestion with page preservation and OCR for photos."""

from __future__ import annotations

import io
import shutil
from pathlib import Path

from models.schemas import DocumentCategory, ExtractedDocument, ExtractedPage
from services.logging_config import logger

MIN_PAGE_CHARS = 40
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".webp", ".bmp"}

TESSERACT_SETUP = """
Tesseract OCR is required to read photos and scanned/photo PDFs.

Windows setup:
1. Install Tesseract from https://github.com/UB-Mannheim/tesseract/wiki
2. During setup, keep the option to add Tesseract to PATH if offered.
3. Typical install path:
     C:\\Program Files\\Tesseract-OCR\\tesseract.exe
4. If it is not on PATH, set TESSERACT_CMD in .env to that full path.
5. Restart this application.

OCR stays on this computer. Images are not sent to a cloud API.
""".strip()


def tesseract_status() -> dict:
    cmd, err = _resolve_tesseract()
    if err:
        return {"available": False, "command": cmd or "", "message": err}
    return {
        "available": True,
        "command": cmd or "tesseract",
        "message": f"Tesseract OCR ready ({cmd or 'tesseract'}).",
    }


def extract_document(
    path: Path,
    original_filename: str,
    category: DocumentCategory,
    sha256: str,
    allow_ocr: bool = True,
) -> ExtractedDocument:
    suffix = path.suffix.lower()
    stored = path.name
    try:
        if suffix == ".pdf":
            return _extract_pdf(path, original_filename, stored, category, sha256, allow_ocr)
        if suffix == ".docx":
            return _extract_docx(path, original_filename, stored, category, sha256)
        if suffix in IMAGE_EXTENSIONS:
            return _extract_image(path, original_filename, stored, category, sha256)
        return ExtractedDocument(
            original_filename=original_filename,
            stored_filename=stored,
            file_type=suffix,
            category=category,
            sha256=sha256,
            extraction_status="error",
            extraction_error="Unsupported file type. Use PDF, DOCX, PNG, JPG, TIF, WEBP or BMP.",
            empty=True,
        )
    except Exception as exc:  # noqa: BLE001 — surface extraction failure to UI
        logger.info("Extraction failed for %s: %s", stored, type(exc).__name__)
        return ExtractedDocument(
            original_filename=original_filename,
            stored_filename=stored,
            file_type=suffix,
            category=category,
            sha256=sha256,
            extraction_status="error",
            extraction_error=f"Extraction failed: {type(exc).__name__}",
            empty=True,
        )


def merge_extracted(
    documents: list[ExtractedDocument], original_filename: str | None = None
) -> ExtractedDocument:
    """Combine several photo pages or files into one question-paper document."""
    if not documents:
        raise ValueError("No documents to merge.")
    if len(documents) == 1:
        return documents[0]
    pages: list[ExtractedPage] = []
    headings: list[str] = []
    page_no = 1
    ocr_used = False
    is_scanned = False
    errors: list[str] = []
    for doc in documents:
        ocr_used = ocr_used or doc.ocr_used
        is_scanned = is_scanned or doc.is_scanned
        headings.extend(doc.headings)
        if doc.extraction_error:
            errors.append(f"{doc.original_filename}: {doc.extraction_error}")
        for page in doc.pages:
            pages.append(
                ExtractedPage(page_number=page_no, text=page.text, heading=page.heading)
            )
            page_no += 1
    full_text = "\n\n".join(p.text for p in pages if p.text).strip()
    empty = len(full_text) < 20
    name = original_filename or " + ".join(d.original_filename for d in documents)
    status = "ok"
    error = None
    if empty:
        status = "empty"
        error = "No readable text could be extracted from the uploaded pages. " + (
            errors[0] if errors else TESSERACT_SETUP
        )
    elif errors:
        status = "ok"
        error = "Some pages had OCR warnings: " + " | ".join(errors[:3])
    return ExtractedDocument(
        original_filename=name,
        stored_filename=documents[0].stored_filename,
        file_type=documents[0].file_type,
        category=documents[0].category,
        page_count=len(pages),
        pages=pages,
        full_text=full_text,
        headings=headings,
        is_scanned=is_scanned,
        ocr_used=ocr_used,
        extraction_status=status,
        extraction_error=error,
        sha256=documents[0].sha256,
        empty=empty,
    )


def _extract_pdf(
    path: Path,
    original_filename: str,
    stored: str,
    category: DocumentCategory,
    sha256: str,
    allow_ocr: bool,
) -> ExtractedDocument:
    import fitz

    doc = fitz.open(path)
    pages: list[ExtractedPage] = []
    headings: list[str] = []
    low_text_pages = 0
    native: list[str] = []
    for i, page in enumerate(doc, start=1):
        text = (page.get_text("text") or "").strip()
        native.append(text)
        if len(text) < MIN_PAGE_CHARS:
            low_text_pages += 1
        heading = _first_heading(text)
        if heading:
            headings.append(heading)
        pages.append(ExtractedPage(page_number=i, text=text, heading=heading))
    page_count = len(doc)
    is_scanned = page_count > 0 and (low_text_pages / page_count) >= 0.5
    ocr_used = False
    ocr_error = None
    if allow_ocr and is_scanned:
        ocr_pages, ocr_ok, ocr_error = _ocr_pdf_pages(doc, native)
        if ocr_ok and ocr_pages:
            pages = ocr_pages
            ocr_used = True
            headings = [h for p in pages if (h := p.heading)]
            is_scanned = False
    doc.close()
    full_text = "\n\n".join(p.text for p in pages if p.text).strip()
    empty = len(full_text) < 20
    status = "ok"
    error = None
    if empty and is_scanned:
        status = "scanned_unreadable"
        error = ocr_error or (
            "This PDF looks like a photo or scan with almost no extractable text. "
            + TESSERACT_SETUP
        )
    elif empty:
        status = "empty"
        error = "No readable text could be extracted."
    elif ocr_used:
        status = "ok"
    return ExtractedDocument(
        original_filename=original_filename,
        stored_filename=stored,
        file_type=".pdf",
        category=category,
        page_count=page_count,
        pages=pages,
        full_text=full_text,
        headings=headings,
        is_scanned=is_scanned and empty,
        ocr_used=ocr_used,
        extraction_status=status,
        extraction_error=error,
        sha256=sha256,
        empty=empty,
    )


def _extract_image(
    path: Path,
    original_filename: str,
    stored: str,
    category: DocumentCategory,
    sha256: str,
) -> ExtractedDocument:
    from PIL import Image, ImageOps, UnidentifiedImageError

    try:
        with Image.open(path) as raw:
            image = ImageOps.exif_transpose(raw)
            image = image.copy()
    except UnidentifiedImageError:
        return ExtractedDocument(
            original_filename=original_filename,
            stored_filename=stored,
            file_type=path.suffix.lower(),
            category=category,
            sha256=sha256,
            extraction_status="error",
            extraction_error="The file could not be opened as an image.",
            empty=True,
        )
    text, ok, ocr_error = _ocr_pil_image(image)
    heading = _first_heading(text)
    empty = len(text.strip()) < 20
    if not ok:
        status = "ocr_unavailable"
        error = ocr_error or TESSERACT_SETUP
    elif empty:
        status = "empty"
        error = (
            "OCR ran but found almost no text. Use a clearer, upright photo with good lighting."
        )
    else:
        status = "ok"
        error = None
    return ExtractedDocument(
        original_filename=original_filename,
        stored_filename=stored,
        file_type=path.suffix.lower(),
        category=category,
        page_count=1,
        pages=[ExtractedPage(page_number=1, text=text, heading=heading)],
        full_text=text,
        headings=[heading] if heading else [],
        is_scanned=True,
        ocr_used=ok,
        extraction_status=status,
        extraction_error=error,
        sha256=sha256,
        empty=empty or not ok,
    )


def _ocr_pdf_pages(doc, native_texts: list[str]) -> tuple[list[ExtractedPage], bool, str | None]:
    from PIL import Image

    pages: list[ExtractedPage] = []
    any_ocr = False
    last_error = None
    for i, page in enumerate(doc, start=1):
        native = native_texts[i - 1] if i - 1 < len(native_texts) else ""
        if len(native) >= MIN_PAGE_CHARS:
            pages.append(
                ExtractedPage(page_number=i, text=native, heading=_first_heading(native))
            )
            continue
        try:
            pix = page.get_pixmap(dpi=200)
            image = Image.open(io.BytesIO(pix.tobytes("png")))
        except Exception as exc:  # noqa: BLE001
            last_error = f"Could not render PDF page {i}: {type(exc).__name__}"
            pages.append(ExtractedPage(page_number=i, text=native, heading=_first_heading(native)))
            continue
        text, ok, err = _ocr_pil_image(image)
        if not ok:
            last_error = err
            pages.append(ExtractedPage(page_number=i, text=native, heading=_first_heading(native)))
            continue
        any_ocr = True
        combined = text.strip() or native
        pages.append(
            ExtractedPage(page_number=i, text=combined, heading=_first_heading(combined))
        )
    return pages, any_ocr, last_error


def _ocr_pil_image(image) -> tuple[str, bool, str | None]:
    cmd, err = _resolve_tesseract()
    if err:
        return "", False, err
    try:
        import pytesseract
        from PIL import ImageOps

        if cmd:
            pytesseract.pytesseract.tesseract_cmd = cmd
        prepared = ImageOps.exif_transpose(image)
        if prepared.mode not in {"L", "RGB"}:
            prepared = prepared.convert("RGB")
        prepared = ImageOps.autocontrast(prepared.convert("L"))
        text = pytesseract.image_to_string(prepared) or ""
        return text.strip(), True, None
    except Exception as exc:  # noqa: BLE001
        logger.info("OCR failed: %s", type(exc).__name__)
        return "", False, f"OCR failed ({type(exc).__name__}). " + TESSERACT_SETUP


def _resolve_tesseract() -> tuple[str | None, str | None]:
    import config as app_config

    candidates: list[str] = []
    if getattr(app_config, "TESSERACT_CMD", ""):
        candidates.append(app_config.TESSERACT_CMD)
    which = shutil.which("tesseract")
    if which:
        candidates.append(which)
    candidates.extend(
        [
            r"C:\Program Files\Tesseract-OCR\tesseract.exe",
            r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe",
        ]
    )
    try:
        import pytesseract
    except ImportError:
        return None, "pytesseract is not installed in the virtual environment. " + TESSERACT_SETUP

    seen: set[str] = set()
    for cmd in candidates:
        if not cmd or cmd in seen:
            continue
        seen.add(cmd)
        path = Path(cmd)
        if path.exists() or cmd == which:
            pytesseract.pytesseract.tesseract_cmd = cmd
            try:
                pytesseract.get_tesseract_version()
                return cmd, None
            except Exception:
                continue
    try:
        pytesseract.get_tesseract_version()
        return "tesseract", None
    except Exception:
        return None, TESSERACT_SETUP


def _extract_docx(
    path: Path,
    original_filename: str,
    stored: str,
    category: DocumentCategory,
    sha256: str,
) -> ExtractedDocument:
    from docx import Document

    document = Document(str(path))
    headings: list[str] = []
    blocks: list[str] = []
    current_heading = None
    for para in document.paragraphs:
        style = (para.style.name if para.style is not None else "") or ""
        text = (para.text or "").strip()
        if not text:
            continue
        if "Heading" in style:
            headings.append(text)
            current_heading = text
            blocks.append(f"## {text}")
        else:
            blocks.append(text)
    full_text = "\n".join(blocks).strip()
    empty = len(full_text) < 20
    pages = [ExtractedPage(page_number=1, text=full_text, heading=current_heading)]
    return ExtractedDocument(
        original_filename=original_filename,
        stored_filename=stored,
        file_type=".docx",
        category=category,
        page_count=1,
        pages=pages,
        full_text=full_text,
        headings=headings,
        extraction_status="empty" if empty else "ok",
        extraction_error="No readable text could be extracted." if empty else None,
        sha256=sha256,
        empty=empty,
    )


def _first_heading(text: str) -> str | None:
    for line in (text or "").splitlines():
        stripped = line.strip()
        if 3 <= len(stripped) <= 80 and stripped.isupper():
            return stripped
        if stripped.lower().startswith(("unit ", "module ", "chapter ", "part ")):
            return stripped
    return None


def preview_text(text: str, limit: int = 800) -> str:
    text = text or ""
    return text[:limit] + ("…" if len(text) > limit else "")
