"""DOCX, PDF and CSV export without secrets or filesystem paths."""

from __future__ import annotations

import csv
import io
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

import config
from models.schemas import ApprovalStatus, EvaluationScheme, QuestionPaper


def _label(status: ApprovalStatus | str) -> str:
    value = status.value if isinstance(status, ApprovalStatus) else str(status)
    if value == ApprovalStatus.APPROVED.value:
        return "FACULTY-APPROVED"
    return "DRAFT — NOT FACULTY-APPROVED"


def export_docx(
    paper: QuestionPaper, scheme: EvaluationScheme, destination: Path | None = None
) -> Path:
    doc = Document()
    title = doc.add_paragraph()
    run = title.add_run("Evaluation Scheme and Solutions")
    run.bold = True
    run.font.size = Pt(18)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    banner = doc.add_paragraph(_label(scheme.approval_status))
    banner.alignment = WD_ALIGN_PARAGRAPH.CENTER
    banner.runs[0].font.color.rgb = RGBColor(0x8A, 0x1C, 0x1C)
    banner.runs[0].bold = True

    doc.add_paragraph(f"Subject: {paper.subject_name} ({paper.subject_code})")
    doc.add_paragraph(f"Job ID: {scheme.job_id}")
    doc.add_paragraph(f"Generated: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')}")
    if scheme.paper_total_marks is not None:
        doc.add_paragraph(f"Printed total marks: {scheme.paper_total_marks}")
    doc.add_paragraph(f"Attempt total (choice-aware): {scheme.attempt_total_marks}")
    doc.add_paragraph(
        "Confidential examination material. Store and share only through authorised college channels."
    )

    if scheme.issues:
        doc.add_heading("Validation notes", level=1)
        for issue in scheme.issues[:30]:
            doc.add_paragraph(f"[{issue.severity}] {issue.message}", style="List Bullet")

    for q in scheme.questions:
        doc.add_heading(f"Question {q.display_number}  ({q.marks if q.marks is not None else '—'} marks)", level=1)
        doc.add_paragraph(q.question_text)
        status = q.syllabus.status.value if q.syllabus else "UNCERTAIN"
        unit = q.syllabus.unit if q.syllabus else ""
        doc.add_paragraph(f"Syllabus status: {status}" + (f" — {unit}" if unit else ""))
        if q.syllabus and q.syllabus.reason:
            doc.add_paragraph(f"Syllabus note: {q.syllabus.reason}")
        if q.syllabus:
            for ev in q.syllabus.evidence[:3]:
                page = f", p.{ev.page_number}" if ev.page_number else ""
                doc.add_paragraph(f"Syllabus evidence: {ev.filename}{page}")
        doc.add_heading("Solution", level=2)
        doc.add_paragraph(q.solution.solution_text if q.solution else "(missing)")
        if q.solution and q.solution.key_points:
            doc.add_paragraph("Key points:")
            for kp in q.solution.key_points:
                doc.add_paragraph(kp, style="List Bullet")
        doc.add_heading("Marking scheme", level=2)
        table = doc.add_table(rows=1, cols=4)
        hdr = table.rows[0].cells
        hdr[0].text = "Component"
        hdr[1].text = "Expected"
        hdr[2].text = "Marks"
        hdr[3].text = "Partial credit / alternatives"
        for c in q.marking_criteria:
            row = table.add_row().cells
            row[0].text = c.step_label
            row[1].text = c.expected_answer
            row[2].text = str(c.marks)
            row[3].text = " | ".join(x for x in [c.partial_credit, c.alternatives] if x)
        doc.add_paragraph(
            f"Allocated total: {q.allocated_marks_sum}  |  Match: {'yes' if q.marks_match else 'NO'}"
        )
        if q.solution:
            for ev in q.solution.evidence[:5]:
                if ev.kind.value == "insufficient":
                    doc.add_paragraph("Evidence: insufficient retrieved support.")
                    continue
                page = f", p.{ev.page_number}" if ev.page_number else ""
                doc.add_paragraph(f"Source: {ev.filename}{page}")
        review = "Approved" if q.approved else "Not approved"
        doc.add_paragraph(f"Review status: {review}")
        if q.faculty_comment:
            doc.add_paragraph(f"Faculty comment: {q.faculty_comment}")
        for w in q.validation_issues:
            doc.add_paragraph(f"Warning: {w}")

    dest = destination or _output_path(scheme.job_id, "scheme.docx")
    dest.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(dest))
    return dest


def export_pdf(
    paper: QuestionPaper, scheme: EvaluationScheme, destination: Path | None = None
) -> Path:
    dest = destination or _output_path(scheme.job_id, "scheme.pdf")
    dest.parent.mkdir(parents=True, exist_ok=True)
    styles = getSampleStyleSheet()
    styles.add(
        ParagraphStyle(name="Banner", fontSize=12, textColor=colors.HexColor("#8A1C1C"), alignment=1, spaceAfter=10)
    )
    styles.add(ParagraphStyle(name="Small", fontSize=9, leading=12))
    story = [
        Paragraph("Evaluation Scheme and Solutions", styles["Title"]),
        Paragraph(_label(scheme.approval_status), styles["Banner"]),
        Paragraph(f"Subject: {paper.subject_name} ({paper.subject_code})", styles["Normal"]),
        Paragraph(f"Job ID: {scheme.job_id}", styles["Small"]),
        Paragraph(
            f"Printed total: {scheme.paper_total_marks} | Attempt total: {scheme.attempt_total_marks}",
            styles["Small"],
        ),
        Spacer(1, 8),
    ]
    for q in scheme.questions:
        story.append(Paragraph(f"Question {q.display_number} ({q.marks} marks)", styles["Heading2"]))
        story.append(Paragraph(_esc(q.question_text), styles["Normal"]))
        if q.syllabus:
            story.append(
                Paragraph(
                    f"Syllabus: {q.syllabus.status.value} — {_esc(q.syllabus.reason)}",
                    styles["Small"],
                )
            )
        if q.solution:
            story.append(Paragraph("<b>Solution</b>", styles["Normal"]))
            story.append(Paragraph(_esc(q.solution.solution_text).replace("\n", "<br/>"), styles["Small"]))
        if q.marking_criteria:
            data = [["Component", "Expected", "Marks", "Partial / alt."]]
            for c in q.marking_criteria:
                data.append(
                    [
                        Paragraph(_esc(c.step_label), styles["Small"]),
                        Paragraph(_esc(c.expected_answer), styles["Small"]),
                        str(c.marks),
                        Paragraph(_esc(" | ".join(x for x in [c.partial_credit, c.alternatives] if x)), styles["Small"]),
                    ]
                )
            table = Table(data, colWidths=[30 * mm, 70 * mm, 18 * mm, 55 * mm])
            table.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1e4d7b")),
                        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                        ("GRID", (0, 0), (-1, -1), 0.3, colors.grey),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("FONTSIZE", (0, 0), (-1, -1), 8),
                    ]
                )
            )
            story.append(table)
        story.append(
            Paragraph(
                f"Allocated {q.allocated_marks_sum} | match={'yes' if q.marks_match else 'NO'} | "
                f"{'Approved' if q.approved else 'Not approved'}",
                styles["Small"],
            )
        )
        story.append(Spacer(1, 8))
    SimpleDocTemplate(
        str(dest),
        pagesize=A4,
        title=f"Evaluation scheme {paper.subject_code}",
        author="Local EvalScheme Generator",
    ).build(story)
    return dest


def export_csv(scheme: EvaluationScheme, destination: Path | None = None) -> Path:
    dest = destination or _output_path(scheme.job_id, "scheme.csv")
    dest.parent.mkdir(parents=True, exist_ok=True)
    with dest.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.writer(fh)
        writer.writerow(
            [
                "question",
                "marks",
                "syllabus_status",
                "allocated_sum",
                "marks_match",
                "approved",
                "review_status",
                "warnings",
            ]
        )
        for q in scheme.questions:
            writer.writerow(
                [
                    q.display_number,
                    q.marks,
                    q.syllabus.status.value if q.syllabus else "",
                    q.allocated_marks_sum,
                    q.marks_match,
                    q.approved,
                    q.status,
                    " | ".join(q.validation_issues),
                ]
            )
    return dest


def export_bytes_docx(paper: QuestionPaper, scheme: EvaluationScheme) -> bytes:
    path = export_docx(paper, scheme)
    return path.read_bytes()


def export_training_dataset(scheme: EvaluationScheme, only_approved: bool = True) -> list[dict]:
    rows = []
    for q in scheme.questions:
        if only_approved and not q.approved:
            continue
        original = q.original_generated or {}
        rows.append(
            {
                "question_text": q.question_text,
                "marks": q.marks,
                "syllabus_classification": q.syllabus.status.value if q.syllabus else None,
                "syllabus_passages": [
                    {"filename": e.filename, "page": e.page_number, "quote": e.quote}
                    for e in (q.syllabus.evidence if q.syllabus else [])
                ],
                "retrieved_passages": [
                    {"filename": e.filename, "page": e.page_number, "quote": e.quote}
                    for e in (q.solution.evidence if q.solution else [])
                ],
                "approved_answer": q.solution.solution_text if q.solution else "",
                "approved_marking_criteria": [c.model_dump() for c in q.marking_criteria],
                "faculty_comment": q.faculty_comment,
                "model_original": original.get("solution"),
                "quality_label": "faculty_approved" if q.approved else "unreviewed",
            }
        )
    return rows


def _output_path(job_id: str, name: str) -> Path:
    folder = config.OUTPUT_DIR / job_id
    folder.mkdir(parents=True, exist_ok=True)
    return folder / name


def _esc(text: str) -> str:
    return (
        (text or "")
        .replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
    )
